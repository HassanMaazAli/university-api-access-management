import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_callout_box(doc, title, text, border_color="1E3A8A", bg_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{title}\n")
    r1.font.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r2 = p.add_run(text)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_placeholder(doc, fig_number, title, description, height_in_inches=3.0):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)
    
    # Dashed border around image space
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(30)
    
    r_icon = p.add_run("[ SCREENSHOT FRAME ]\n")
    r_icon.font.size = Pt(28)
    
    r_tag = p.add_run(f"[ PASTE SCREENSHOT HERE ]\n")
    r_tag.font.bold = True
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(12)
    r_tag.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    r_desc = p.add_run(f"{description}\n(Recommended size: Full Width or ~6.5 inches)")
    r_desc.font.italic = True
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Caption below table
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(f"Figure {fig_number}: {title}")
    r_cap.font.bold = True
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(10)
    r_cap.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def build_srs_document(output_path):
    doc = docx.Document()
    
    # Page Margins: 1 inch everywhere
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # ==========================================
    # COVER / HEADER SECTION
    # ==========================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("AIR UNIVERSITY ISLAMABAD\n")
    r_inst.font.bold = True
    r_inst.font.size = Pt(16)
    r_inst.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
    
    r_camp = p_inst.add_run("Kamra Campus — Department of Computer Science\n")
    r_camp.font.bold = True
    r_camp.font.size = Pt(13)
    r_camp.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_sub = p_inst.add_run("CS-301L: Full Stack Development | Lab Assignment Report & SRS\n")
    r_sub.font.bold = True
    r_sub.font.size = Pt(11.5)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Metadata Grid Table
    meta_table = doc.add_table(rows=4, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Course Name:", True), ("Full Stack Development", False), ("Date:", True), ("10/20/2026", False)],
        [("Course Code:", True), ("CS-301L", False), ("Teacher / Instructor:", True), ("M. Fahad Qasim", False)],
        [("Lab Number:", True), ("Lab 07 (20 Marks)", False), ("Semester / Section:", True), ("BSCS VB", False)],
        [("Discipline:", True), ("Computer Science", False), ("Project Title:", True), ("University API Access System", False)],
    ]
    
    for r_idx, row in enumerate(meta_data):
        for c_idx, (text, is_label) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            if is_label:
                set_cell_background(cell, "F1F5F9")
            else:
                set_cell_background(cell, "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            if is_label:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Submission Card
    sub_table = doc.add_table(rows=1, cols=2)
    sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1 = sub_table.cell(0, 0)
    c2 = sub_table.cell(0, 1)
    set_cell_background(c1, "EFF6FF")
    set_cell_background(c2, "EFF6FF")
    set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
    set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
    
    p_s1 = c1.paragraphs[0]
    p_s1.paragraph_format.space_after = Pt(0)
    p_s1.add_run("Student Name: ").font.bold = True
    p_s1.add_run("[ Enter Your Full Name ]")
    
    p_s2 = c2.paragraphs[0]
    p_s2.paragraph_format.space_after = Pt(0)
    p_s2.add_run("Registration / Roll No: ").font.bold = True
    p_s2.add_run("[ Enter Your Roll Number ]")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # CLO Mapping Callout
    create_callout_box(
        doc,
        "Course Learning Outcomes (CLO) Mapping",
        "• CLO-3: Design and develop a responsive frontend with reusable components and AJAX-based dynamic updates.\n"
        "• CLO-4: Analyze and implement server-side operations, including API endpoints, database access, validation, and CRUD functionality."
    )

    # ==========================================
    # SECTION 1: SYSTEM OVERVIEW & SCOPE
    # ==========================================
    h1 = doc.add_heading("1. Executive Summary & System Overview", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    
    doc.add_paragraph(
        "The University API Access Management System is an enterprise-grade full-stack web application designed "
        "to manage secure academic data interchange across administrators, faculty members, and students. "
        "Unlike conventional session-cookie architectures, this system implements header-based cryptographic API key "
        "authentication, strict role-based access control (RBAC), and automated rate limiting to protect university records."
    )
    
    doc.add_paragraph(
        "The project strictly complies with the specified constraints:\n"
        "1. Backend: Node.js and Express.js RESTful service with modular MVC architecture.\n"
        "2. Database: MySQL/MariaDB utilizing parameterized queries via mysql2 connection pooling.\n"
        "3. Frontend: HTML5, CSS3, JavaScript (ES6), Bootstrap 5 SaaS dashboard, and jQuery AJAX.\n"
        "4. Authentication: Header-only API keys (x-api-key) with SHA-256 one-way cryptographic hashing.\n"
        "5. Strict Data Isolation: Student identity is derived exclusively from the verified API key hash. "
        "Client-provided student IDs are completely disregarded, eliminating Insecure Direct Object References (IDOR)."
    )

    # Architecture Overview Table
    tech_table = doc.add_table(rows=6, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer", "Technologies Used", "Architectural Role & Functionality"]
    for i, h in enumerate(headers):
        cell = tech_table.cell(0, i)
        set_cell_background(cell, "002060")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    rows_data = [
        ("Presentation Layer", "HTML5, CSS3, Bootstrap 5.3, Bootstrap Icons, Chart.js", "Responsive SaaS dashboard with dark navy navigation, KPI cards, and modals."),
        ("Dynamic Client Layer", "jQuery 3.7.1 AJAX (ApiClient wrapper)", "Asynchronous CRUD operations, header injection (x-api-key), and toast alerts."),
        ("Application Layer", "Node.js (v24), Express.js (v4.21), Helmet, Morgan, CORS", "REST API routing, security middleware, centralized error handling pipeline."),
        ("Security & Auth Layer", "Node.js Crypto (SHA-256), Custom RBAC Middleware", "Cryptographic key generation, SHA-256 hash lookup, and atomic quota enforcement."),
        ("Persistence Layer", "MySQL 5.7+ / MariaDB 10.4+ (mysql2/promise pool)", "Relational database with foreign keys, indexes, cascades, and audit logs.")
    ]

    for r_idx, row in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tech_table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 != 0 else "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 2: DATABASE DESIGN & SCHEMA
    # ==========================================
    h2 = doc.add_heading("2. Database Design & Relational Schema", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The relational database university_api_system comprises five normalized tables with strict foreign key "
        "constraints, indexed lookups, and audit trails."
    )

    schema_table = doc.add_table(rows=6, cols=4)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Table Name", "Primary Key", "Foreign Keys", "Key Fields & Constraints"]
    for i, h in enumerate(s_headers):
        cell = schema_table.cell(0, i)
        set_cell_background(cell, "002060")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    s_rows = [
        ("students", "id (INT AUTO_INCREMENT)", "None", "student_number (UNIQUE), full_name, email (UNIQUE), department, semester, created_at"),
        ("teachers", "id (INT AUTO_INCREMENT)", "None", "teacher_number (UNIQUE), full_name, email (UNIQUE), department, created_at"),
        ("academic_marks", "id (INT AUTO_INCREMENT)", "student_id -> students(id)\nteacher_id -> teachers(id)", "subject_name, total_marks (100), obtained_marks (DECIMAL), semester. Cascading deletes."),
        ("api_keys", "id (INT AUTO_INCREMENT)", "Polymorphic owner_id", "key_prefix (VARCHAR 24), key_hash (VARCHAR 128 UNIQUE), owner_type ('student','teacher','admin'), request_limit, request_count, status ('active','revoked')"),
        ("api_request_logs", "id (INT AUTO_INCREMENT)", "api_key_id -> api_keys(id)", "endpoint, http_method, response_status, requested_at. Indexed on timestamp and status.")
    ]

    for r_idx, row in enumerate(s_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = schema_table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 != 0 else "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9.0)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ER Diagram Placeholder
    add_image_placeholder(
        doc,
        fig_number="2.1",
        title="Entity-Relationship Diagram (ERD) of university_api_system",
        description="Insert the database ER diagram generated from phpMyAdmin Designer, MySQL Workbench, or Mermaid visualization."
    )

    # ==========================================
    # SECTION 3: SYSTEM ARCHITECTURE & SECURITY
    # ==========================================
    h3 = doc.add_heading("3. Security Architecture & Access Control Specification", level=1)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The security model addresses all 8 lab requirements through multi-layered defensive controls:"
    )

    doc.add_paragraph(
        "• Cryptographic Key Generation & Hashing: Keys are minted as uap_<role>_<random_24_bytes_hex>. "
        "The raw key is returned exactly once to the client upon creation. The server hashes the key with SHA-256 "
        "(crypto.createHash('sha256')) and stores only the 64-character hash and a safe display prefix (e.g., uap_student_a1b2...).\n\n"
        "• Elimination of URL Key Leakage: Keys are transmitted solely in the x-api-key HTTP header. "
        "Any request with ?api_key= or ?key= in the query string is immediately intercepted and rejected with HTTP 400 Bad Request.\n\n"
        "• Key-Derived Identity & Anti-IDOR: In studentController.js, student identity is obtained from req.auth.ownerId "
        "derived from the verified API key hash. Client-supplied student IDs in URLs, bodies, or headers are ignored, preventing "
        "cross-account horizontal privilege escalation.\n\n"
        "• Atomic Concurrency Rate Limiter: When requests arrive, MySQL executes an atomic parameterized update:\n"
        "  UPDATE api_keys SET request_count = request_count + 1, last_used_at = NOW() WHERE id = ? AND request_count < request_limit;\n"
        "If affectedRows is 0, the quota has been exhausted; the server immediately aborts processing and issues HTTP 429 Too Many Requests.\n\n"
        "• Real-Time Key Revocation: Admins can toggle key status to 'revoked'. Subsequent requests using that key "
        "immediately yield HTTP 403 Forbidden without database modifications."
    )

    # ==========================================
    # SECTION 4: REST API ENDPOINTS SPECIFICATION
    # ==========================================
    h4 = doc.add_heading("4. REST API Endpoints Specification", level=1)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)

    api_table = doc.add_table(rows=13, cols=4)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Endpoint", "Method", "Auth & Role", "Description & Response Status"]
    for i, h in enumerate(api_headers):
        cell = api_table.cell(0, i)
        set_cell_background(cell, "002060")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    api_rows = [
        ("/api/auth/validate", "POST", "Public / Valid Key", "Validates key and returns safe role & account info. HTTP 200 / 401 / 403"),
        ("/api/auth/me", "GET", "Any Valid Key", "Returns authenticated caller identity attached to req.auth. HTTP 200 / 401"),
        ("/api/admin/stats", "GET", "Admin Only", "Returns total students, faculty, active/revoked keys, request count. HTTP 200 / 403"),
        ("/api/admin/students", "GET", "Admin Only", "Searchable student directory with department and semester filters. HTTP 200 / 403"),
        ("/api/admin/teachers", "GET", "Admin Only", "Searchable faculty directory with department filters. HTTP 200 / 403"),
        ("/api/admin/api-keys", "GET", "Admin Only", "Lists all keys with safe prefixes, status, and usage. Excludes hashes. HTTP 200 / 403"),
        ("/api/admin/api-keys", "POST", "Admin Only", "Generates new API key with custom request quota limit. HTTP 201 Created / 400"),
        ("/api/admin/api-keys/:id/limit", "PATCH", "Admin Only", "Updates API request quota limit. Preserves existing usage. HTTP 200 / 400"),
        ("/api/admin/api-keys/:id/revoke", "PATCH", "Admin Only", "Immediately revokes key, rendering it inactive. HTTP 200 / 404"),
        ("/api/admin/logs", "GET", "Admin Only", "Paginated and filterable audit request log register. HTTP 200 / 403"),
        ("/api/student/profile", "GET", "Student Only", "Returns authenticated student's profile (derived from key). HTTP 200 / 403"),
        ("/api/student/marks", "GET", "Student Only", "Returns student's isolated marks, percentage, and grade. HTTP 200 / 403"),
    ]

    for r_idx, row in enumerate(api_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = api_table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 != 0 else "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9.0)
            if c_idx == 1:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 5: IMPLEMENTATION CODE EXCERPTS
    # ==========================================
    h5 = doc.add_heading("5. Implementation Details & Core Code Excerpts", level=1)
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(6)

    def add_code_block(doc, title, code):
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(f"Listing: {title}")
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
        
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(code)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_code_block(
        doc,
        "SHA-256 Key Hashing & Cryptographic Generation (services/keyService.js)",
        "function hashKey(rawKey) {\n"
        "  if (!rawKey || typeof rawKey !== 'string') return '';\n"
        "  return crypto.createHash('sha256').update(rawKey.trim()).digest('hex');\n"
        "}\n\n"
        "function generateRawKey(role = 'live') {\n"
        "  const randomBytes = crypto.randomBytes(24).toString('hex');\n"
        "  return `uap_${role}_${randomBytes}`;\n"
        "}"
    )

    add_code_block(
        doc,
        "Header Extraction, Quota Enforcement & Atomic Increment (middleware/authMiddleware.js)",
        "// Reject key passed in query string\n"
        "if (req.query.api_key || req.query.key) {\n"
        "  return res.status(400).json({ error: 'Bad Request', message: 'API keys must be in x-api-key header.' });\n"
        "}\n\n"
        "const rawKey = req.headers['x-api-key'];\n"
        "if (!rawKey) return res.status(401).json({ error: 'Unauthorized', message: 'Missing x-api-key header.' });\n\n"
        "// Check rate limit\n"
        "if (keyData.request_count >= keyData.request_limit) {\n"
        "  return res.status(429).json({ error: 'Too Many Requests', message: 'API request quota exceeded.' });\n"
        "}\n\n"
        "// Concurrency-safe atomic increment\n"
        "const [updateResult] = await pool.execute(\n"
        "  `UPDATE api_keys SET request_count = request_count + 1, last_used_at = NOW() ` +\n"
        "  `WHERE id = ? AND request_count < request_limit`,\n"
        "  [keyData.id]\n"
        ");\n"
        "if (updateResult.affectedRows === 0) {\n"
        "  return res.status(429).json({ error: 'Too Many Requests', message: 'Rate limit reached.' });\n"
        "}"
    )

    add_code_block(
        doc,
        "Key-Derived Student Marks Query (controllers/studentController.js)",
        "// Server derives student ID strictly from authenticated key metadata\n"
        "const studentId = req.auth.ownerId;\n\n"
        "const [marks] = await pool.execute(\n"
        "  `SELECT m.id, m.subject_name, m.total_marks, m.obtained_marks, m.semester, ` +\n"
        "  `t.full_name AS teacher_name ` +\n"
        "  `FROM academic_marks m ` +\n"
        "  `JOIN teachers t ON m.teacher_id = t.id ` +\n"
        "  `WHERE m.student_id = ? ORDER BY m.semester ASC`, [studentId]\n"
        ");"
    )

    # ==========================================
    # SECTION 6: SCREENSHOT CHECKLIST & FRAMED PLACEHOLDERS
    # ==========================================
    h6 = doc.add_heading("6. System Verification Screenshots (Submission Artifacts)", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Paste your operational screenshots into the framed placeholders below. "
        "Each section corresponds directly to the grading rubric criteria for Lab 07."
    )

    # 1. Admin Dashboard Overview
    add_image_placeholder(
        doc,
        fig_number="6.1",
        title="Administrator Dashboard — Overview KPIs & Live Charts",
        description="Navigate to http://localhost:3000/admin. Capture the top stat cards (Total Students, Teachers, Active Keys, Total Requests), the Doughnut Chart, and the Recent API Activity log table."
    )

    # 2. Managed Keys Table
    add_image_placeholder(
        doc,
        fig_number="6.2",
        title="Admin Panel — Managed API Keys Directory & Status Badges",
        description="In /admin, select 'API Keys Management'. Capture the table displaying key names, safe prefixes (e.g., uap_student_...), role badges, usage progress bars, and action buttons."
    )

    # 3. Key Generation Modal
    add_image_placeholder(
        doc,
        fig_number="6.3",
        title="Admin Panel — Cryptographic API Key Generation Dialog",
        description="Click 'Generate New Key' in /admin. Show the filled modal selecting a student/teacher, description, and request limit (e.g., 150 calls)."
    )

    # 4. Key Reveal Single-Time Modal
    add_image_placeholder(
        doc,
        fig_number="6.4",
        title="Admin Panel — Single-Time Key Reveal with Copy to Clipboard",
        description="After generating a key, capture the green modal displaying the plaintext raw key in a dark box with the copy button and security warning."
    )

    # 5. Request Limit Update & Revocation
    add_image_placeholder(
        doc,
        fig_number="6.5",
        title="Admin Panel — Request Limit Update & Revocation Modals",
        description="Capture the modal dialog for modifying request limit quota (PATCH /api/admin/api-keys/:id/limit) and the permanent revocation confirmation dialog."
    )

    # 6. Audit Request Logs
    add_image_placeholder(
        doc,
        fig_number="6.6",
        title="Admin Panel — Request Audit Logs Register with Status Badges",
        description="In /admin, click 'Audit Request Logs'. Capture the table showing recorded endpoints, HTTP methods, timestamps, and status codes (200, 401, 403, 429)."
    )

    # 7. Settings Page
    add_image_placeholder(
        doc,
        fig_number="6.7",
        title="Settings Portal — API Key Connection & AJAX Validation",
        description="Open http://localhost:3000/settings. Capture the key input box, successful validation badge, connected user identity card, and quota allocation bar."
    )

    # 8. Student Dashboard Marks
    add_image_placeholder(
        doc,
        fig_number="6.8",
        title="Student Portal — Key-Isolated Academic Marks Table",
        description="Logged in as Zainab Ahmed at /student/dashboard. Capture the enrolled subjects, total marks, obtained marks, calculated percentage, and assigned faculty instructor."
    )

    # 9. Student Performance Chart
    add_image_placeholder(
        doc,
        fig_number="6.9",
        title="Student Portal — Academic Performance Visualization Chart",
        description="In /student/dashboard, capture the Chart.js bar chart showing subject scores compared against maximum total marks."
    )

    # 10. Teacher Dashboard
    add_image_placeholder(
        doc,
        fig_number="6.10",
        title="Faculty Portal — Multi-Department Academic Marks Register",
        description="Logged in as Dr. Alan Turing at /teacher/dashboard. Capture the overview summary cards, student marks across departments, and the department/semester filters."
    )

    # 11. Security Test: Student Denied Teacher Endpoint
    add_image_placeholder(
        doc,
        fig_number="6.11",
        title="Security Verification — Student Blocked from Faculty Route (HTTP 403 Forbidden)",
        description="Demonstrate role boundary enforcement. Using DevTools Network tab or curl/Postman, show a student key attempting to call GET /api/teacher/students resulting in HTTP 403 Forbidden."
    )

    # 12. Security Test: Quota Exceeded 429
    add_image_placeholder(
        doc,
        fig_number="6.12",
        title="Security Verification — Rate Limit Quota Exceeded (HTTP 429 Too Many Requests)",
        description="Click 'Simulate Exceeded Quota Request' on the home page or send a request with Hamza Tariq's quota-exceeded key. Capture the HTTP 429 toast and JSON response."
    )

    # 13. Security Test: Revoked Key 403
    add_image_placeholder(
        doc,
        fig_number="6.13",
        title="Security Verification — Revoked Key Rejection (HTTP 403 Forbidden)",
        description="Click 'Simulate Revoked Request' on the home page. Capture the HTTP 403 toast notification stating 'This API key has been revoked and is no longer active.'"
    )

    # 14. Automated Test Suite
    add_image_placeholder(
        doc,
        fig_number="6.14",
        title="Automated Test Suite — All 27 End-to-End Tests Passing (npm test)",
        description="Capture your terminal showing the successful execution of 'npm test' with all 27 unit and integration tests passing with 0 failures."
    )

    # ==========================================
    # SECTION 7: DEMO VIDEO SCRIPT
    # ==========================================
    h7 = doc.add_heading("7. Video Demonstration Script (2–3 Minutes)", level=1)
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Use this structured script while recording your lab demonstration video to ensure full marks for the video deliverable:"
    )

    script_table = doc.add_table(rows=6, cols=3)
    script_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sc_headers = ["Timestamp", "Visual Action on Screen", "Verbal Narration Guide"]
    for i, h in enumerate(sc_headers):
        cell = script_table.cell(0, i)
        set_cell_background(cell, "002060")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    script_rows = [
        ("0:00 - 0:30", "Show Home Page (http://localhost:3000) & Architecture", 
         "\"Hello, this is my submission for Lab 07: University API Access Management System. The application is built using Node.js, Express, MySQL, and jQuery AJAX. Authentication is header-based using x-api-key without plaintext storage.\""),
        ("0:30 - 1:00", "Admin Console (/admin): Stats, Generate Key, Revoke", 
         "\"Here in the Admin Console, we see live KPIs. We generate a new key for student Ayesha Khan with a quota of 250 requests. The system reveals the raw key once. Only its SHA-256 hash is saved. We can also modify limits and revoke keys.\""),
        ("1:00 - 1:30", "Settings (/settings) & Student Portal (/student/dashboard)", 
         "\"Next, we navigate to /settings and validate Zainab Ahmed's student key. In the Student Dashboard, Zainab sees only her own profile and marks. The server derives her identity strictly from the key hash, preventing ID tampering.\""),
        ("1:30 - 2:00", "Faculty Dashboard (/teacher/dashboard)", 
         "\"Switching to Dr. Alan Turing's faculty key, we access the Teacher Portal. Dr. Turing can view and filter academic marks across all students and departments, but cannot perform administrative functions.\""),
        ("2:00 - 2:30", "Rate Limit (429) & Revoked Key (403) Demonstrations", 
         "\"Finally, testing our quota-exceeded key immediately yields HTTP 429 Too Many Requests via atomic SQL counters. Testing our revoked key yields HTTP 403 Forbidden. All requests are recorded in the audit log. Thank you.\"")
    ]

    for r_idx, row in enumerate(script_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = script_table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 != 0 else "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9.0)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ==========================================
    # SECTION 8: CLO REFLECTIONS & CONCLUSION
    # ==========================================
    h8 = doc.add_heading("8. Conclusion & CLO Learning Reflections", level=1)
    h8.paragraph_format.space_before = Pt(12)
    h8.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "1. CLO-3 (Frontend Engineering): Developed a clean, responsive SaaS administration interface using Bootstrap 5, "
        "reusable custom CSS variables, and modular jQuery AJAX helpers. All data interactions happen asynchronously without page "
        "reloads, providing user feedback through Bootstrap Toasts and dynamic Chart.js data visualizations.\n\n"
        "2. CLO-4 (Server-Side Architecture): Designed a secure Express.js REST API integrated with MySQL connection pooling. "
        "Implemented industry-standard API security practices including SHA-256 key hashing, parameterized SQL queries, "
        "atomic rate limiting, and strict key-derived data isolation, fulfilling all requirements of Lab 07."
    )

    doc.save(output_path)
    print(f"SUCCESS: SRS Document successfully generated at: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\FINE LAPTOPS\.gemini\antigravity\scratch\university-api-system"
    out_file = os.path.join(out_dir, "Lab07_SRS_University_API_Access_Management_System.docx")
    build_srs_document(out_file)
